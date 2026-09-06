#!/usr/bin/env python3
"""Generate CodeLens pitch deck (PowerPoint)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# CodeLens palette (matches frontend styles.css)
BG = RGBColor(11, 16, 32)
SURFACE = RGBColor(18, 26, 47)
ACCENT = RGBColor(91, 140, 255)
ACCENT2 = RGBColor(124, 92, 255)
TEXT = RGBColor(238, 242, 255)
MUTED = RGBColor(154, 168, 199)
LOW = RGBColor(6, 214, 160)
MEDIUM = RGBColor(255, 209, 102)
HIGH = RGBColor(255, 107, 107)
WHITE = RGBColor(255, 255, 255)

OUT = Path(__file__).resolve().parents[1] / "docs" / "CodeLens-Presentation.pptx"


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_accent_bar(slide, left=Inches(0), top=Inches(0), width=Inches(0.12), height=Inches(7.5)) -> None:
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_title_block(
    slide,
    title: str,
    subtitle: str = "",
    *,
    title_size: int = 40,
    subtitle_size: int = 18,
    left=Inches(0.85),
    top=Inches(0.55),
    width=Inches(11.5),
) -> None:
    box = slide.shapes.add_textbox(left, top, width, Inches(1.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(title_size)
    p.font.bold = True
    p.font.color.rgb = TEXT
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(subtitle_size)
        p2.font.color.rgb = MUTED
        p2.space_before = Pt(8)


def add_bullets(
    slide,
    items: list[str],
    *,
    left=Inches(0.95),
    top=Inches(1.85),
    width=Inches(11.2),
    height=Inches(5.2),
    font_size: int = 20,
    color: RGBColor = TEXT,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(10)


def add_card(
    slide,
    x,
    y,
    w,
    h,
    title: str,
    body: str,
    accent: RGBColor = ACCENT,
) -> None:
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = SURFACE
    card.line.color.rgb = RGBColor(42, 54, 84)
    card.line.width = Pt(1)

    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, w, Inches(0.08))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()

    tbox = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.22), w - Inches(0.4), Inches(0.5))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    tp.font.bold = True
    tp.font.size = Pt(16)
    tp.font.color.rgb = TEXT

    bbox = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.65), w - Inches(0.4), h - Inches(0.8))
    bf = bbox.text_frame
    bf.word_wrap = True
    bp = bf.paragraphs[0]
    bp.text = body
    bp.font.size = Pt(13)
    bp.font.color.rgb = MUTED


def slide_title_only(prs: Presentation, title: str, subtitle: str = "") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG)
    add_accent_bar(slide)
    add_title_block(slide, title, subtitle)


def slide_bullets(prs: Presentation, title: str, bullets: list[str], subtitle: str = "") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG)
    add_accent_bar(slide)
    add_title_block(slide, title, subtitle)
    add_bullets(slide, bullets)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 — Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG)
    add_accent_bar(slide, width=Inches(0.18))
    badge = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(1.6), Inches(2.4), Inches(0.45)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 45, 80)
    badge.line.fill.background()
    bt = badge.text_frame.paragraphs[0]
    bt.text = "PR Review Intelligence"
    bt.font.size = Pt(12)
    bt.font.bold = True
    bt.font.color.rgb = ACCENT
    bt.alignment = PP_ALIGN.CENTER

    title = slide.shapes.add_textbox(Inches(0.95), Inches(2.2), Inches(11), Inches(1.2))
    tp = title.text_frame.paragraphs[0]
    tp.text = "CodeLens"
    tp.font.size = Pt(54)
    tp.font.bold = True
    tp.font.color.rgb = TEXT

    sub = slide.shapes.add_textbox(Inches(0.95), Inches(3.35), Inches(10.5), Inches(1))
    sp = sub.text_frame.paragraphs[0]
    sp.text = "Understand a pull request before reviewing it."
    sp.font.size = Pt(24)
    sp.font.color.rgb = MUTED

    tag = slide.shapes.add_textbox(Inches(0.95), Inches(5.9), Inches(10), Inches(0.5))
    tagp = tag.text_frame.paragraphs[0]
    tagp.text = "GitHub · Risk analysis · AI summaries · Human review stays in control"
    tagp.font.size = Pt(14)
    tagp.font.color.rgb = ACCENT

    # 2 — Problem
    slide_bullets(
        prs,
        "The problem",
        [
            "Large PRs hide scope, risk, and impact across diffs, commits, and comments.",
            "Reviewers spend time reconstructing context before they can give useful feedback.",
            "Critical paths (auth, payments, security) are easy to miss in a wall of green/red.",
            "Teams need orientation — not another noisy bot with hundreds of generic findings.",
        ],
        subtitle="From “Where do I start?” to focused, explainable review",
    )

    # 3 — Solution
    slide_bullets(
        prs,
        "The CodeLens answer",
        [
            "Connect GitHub → pick a repo → analyze any open pull request.",
            "Get a structured brief: risk score, executive summary, and where to focus first.",
            "Walk file-by-file and commit-by-commit with diffs and optional AI summaries.",
            "Human review remains final — CodeLens informs; you decide approve & merge.",
        ],
        subtitle="What changed · How risky · Where to look",
    )

    # 4 — User journey
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG)
    add_accent_bar(slide)
    add_title_block(slide, "User journey", "End-to-end flow in the dashboard")
    steps = [
        ("1", "Sign in", "GitHub OAuth — repo read access"),
        ("2", "Select repo", "Dashboard lists your repositories"),
        ("3", "Open PR", "View open pull requests"),
        ("4", "Analyze", "One click — rules + optional Groq AI"),
        ("5", "Review", "Tabs: Overview · Changes · Commits · Discussion · Risk"),
        ("6", "Act", "Approve or merge on GitHub from the report (with confirmation)"),
    ]
    for i, (num, title, body) in enumerate(steps):
        col = i % 3
        row = i // 3
        x = Inches(0.95) + col * Inches(4.05)
        y = Inches(2.0) + row * Inches(2.35)
        add_card(slide, x, y, Inches(3.75), Inches(2.05), f"{num}. {title}", body, accent=ACCENT if i % 2 == 0 else ACCENT2)

    # 5 — Report tabs
    slide_bullets(
        prs,
        "PR report — five lenses",
        [
            "Overview — PR summary, executive summary, ranked focus areas (rule-based).",
            "Changes — CodeRabbit-style file walkthrough with full diffs + AI per-file notes.",
            "Commits — Timeline with per-commit file patches (additions/deletions).",
            "Discussion — Summarized review comments and conversation activity.",
            "Risk analysis — Six weighted dimensions with explainable findings + evidence.",
        ],
        subtitle="One page, multiple ways to orient before deep review",
    )

    # 6 — Risk dimensions
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG)
    add_accent_bar(slide)
    add_title_block(slide, "Analysis engine", "Weighted risk score 0–100 → High / Medium / Low")
    dims = [
        ("Change & drift", "Scope broader than PR title? Multi-module clusters?"),
        ("Code volume", "Lines changed vs baseline; oversized files"),
        ("Functionality", "One coherent change vs scattered edits"),
        ("Critical paths", "Auth, payments, permissions, data, infra, API"),
        ("Security", "Secrets, eval, SQL concat, XSS patterns, TLS disable"),
        ("Test coverage", "Source changed without matching test files"),
    ]
    for i, (name, desc) in enumerate(dims):
        col = i % 2
        row = i // 2
        x = Inches(0.95) + col * Inches(6.1)
        y = Inches(1.9) + row * Inches(1.55)
        accent = [HIGH, MEDIUM, LOW, ACCENT, ACCENT2, LOW][i]
        add_card(slide, x, y, Inches(5.85), Inches(1.35), name, desc, accent=accent)

    # 7 — AI layer
    slide_bullets(
        prs,
        "AI layer (optional)",
        [
            "Providers: Groq, OpenAI, Perplexity, OpenRouter — configured via environment.",
            "Executive summary, PR overview, discussion summary, per-file change notes.",
            "Clear badges: AI · groq vs Rules — so reviewers know the source.",
            "Graceful fallback to rule-based text when AI is unavailable.",
            "MCP server exposes analyzers to Cursor IDE for local agent workflows.",
        ],
        subtitle="Explainable signals first; AI augments, not replaces",
    )

    # 8 — GitHub integration
    slide_bullets(
        prs,
        "GitHub integration",
        [
            "OAuth App — login, list repos/PRs, approve & merge via your GitHub token.",
            "GitHub App (production) — webhooks auto-analyze on pull_request events.",
            "Checks API — risk conclusion appears on the PR Checks tab on github.com.",
            "PR comments — post CodeLens summary back to the conversation (optional).",
            "Signed embed URLs — report viewable from Check details / github.com iframe.",
        ],
        subtitle="Standalone web app today · native GitHub surfaces when App is installed",
    )

    # 9 — Architecture
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG)
    add_accent_bar(slide)
    add_title_block(slide, "Architecture", "FastAPI + React + SQLite · Docker & Kubernetes ready")
    arch_items = [
        ("Frontend", "React / Vite — dashboard, PR report, embed route"),
        ("Backend API", "FastAPI — OAuth, analyze pipeline, webhooks, merge/approve"),
        ("Analyzers", "Rule engine + optional Groq — engine.py, summarize.py"),
        ("Database", "SQLite (dev) — cached reports, installations, webhook audit"),
        ("Deploy", "Docker Compose + K8s manifests in deploy/kubernetes/"),
    ]
    for i, (name, desc) in enumerate(arch_items):
        y = Inches(1.85) + i * Inches(1.05)
        pill = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), y, Inches(2.2), Inches(0.55))
        pill.fill.solid()
        pill.fill.fore_color.rgb = RGBColor(30, 45, 80)
        pill.line.fill.background()
        pp = pill.text_frame.paragraphs[0]
        pp.text = name
        pp.font.size = Pt(13)
        pp.font.bold = True
        pp.font.color.rgb = ACCENT
        pp.alignment = PP_ALIGN.CENTER
        desc_box = slide.shapes.add_textbox(Inches(3.35), y + Inches(0.05), Inches(9.2), Inches(0.55))
        dp = desc_box.text_frame.paragraphs[0]
        dp.text = desc
        dp.font.size = Pt(16)
        dp.font.color.rgb = TEXT

    # 10 — Demo
    slide_bullets(
        prs,
        "Live demo checklist",
        [
            "Login with GitHub → Dashboard → pick a repository.",
            "Open an open PR → auto-analyze (or Re-analyze).",
            "Show Overview: risk banner + focus areas.",
            "Changes tab: expand a file diff + AI summary badge.",
            "Risk analysis tab: dimension scores and findings with file evidence.",
            "Optional: Merge into main (confirm modal) — real GitHub API call.",
            "No GitHub? POST /api/demo/analyze with fixture PR JSON.",
        ],
        subtitle="localhost:5173 (dev) · localhost:8080 (Docker)",
    )

    # 11 — Differentiators
    slide_bullets(
        prs,
        "Why CodeLens?",
        [
            "Focused brief — top focus areas, not hundreds of low-value comments.",
            "Explainable — every finding ties to files, metrics, or rule evidence.",
            "Transparent AI — badges show when Groq vs rules generated text.",
            "Full diff story — files + commits + discussion in one report.",
            "Production path — webhooks, Checks, embed, Docker/K8s included.",
        ],
        subtitle="Inspired by CodeRabbit · built for hackathon clarity",
    )

    # 12 — Roadmap
    slide_bullets(
        prs,
        "Roadmap & gaps",
        [
            "Code quality dimension — maintainability, complexity (PRD gap).",
            "Org-level dashboards — risk trends across repos.",
            "PostgreSQL for multi-replica production deploys.",
            "Deeper repo context — blame, history, dependency graphs.",
            "GitHub Marketplace listing for one-click install.",
        ],
        subtitle="Strong prototype today · clear path to production",
    )

    # 13 — Thank you
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG)
    add_accent_bar(slide, width=Inches(0.18))
    center = slide.shapes.add_textbox(Inches(1), Inches(2.4), Inches(11.3), Inches(1.2))
    cp = center.text_frame.paragraphs[0]
    cp.text = "Thank you"
    cp.font.size = Pt(48)
    cp.font.bold = True
    cp.font.color.rgb = TEXT
    cp.alignment = PP_ALIGN.CENTER

    q = slide.shapes.add_textbox(Inches(1), Inches(3.6), Inches(11.3), Inches(1.5))
    qf = q.text_frame
    qf.word_wrap = True
    qp = qf.paragraphs[0]
    qp.text = "Questions?"
    qp.font.size = Pt(28)
    qp.font.color.rgb = ACCENT
    qp.alignment = PP_ALIGN.CENTER
    qp2 = qf.add_paragraph()
    qp2.text = "github.com · CodeLens — Understand before you review"
    qp2.font.size = Pt(16)
    qp2.font.color.rgb = MUTED
    qp2.alignment = PP_ALIGN.CENTER
    qp2.space_before = Pt(16)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
